"""
Parser pour les fichiers DOCX.

Stratégie d'extraction des titres :
1. D'abord, vérifier si le document utilise les styles Heading natifs de Word
2. Sinon, extraire les métadonnées de formatage de chaque paragraphe et
   envoyer au LLM pour détecter la hiérarchie des titres
"""

import logging
from docx import Document
from lxml import etree
from .base_parser import BaseParser, ParsedDocument, Section
from .llm_heading_detector import detect_headings_with_llm

# Namespace XML de Word
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
W = '{' + W_NS + '}'
NS = {'w': W_NS}

logger = logging.getLogger('mini_fiche.parser.docx')


class DocxParser(BaseParser):
    """Parser spécialisé pour les fichiers .docx"""

    def parse(self, file_path: str) -> ParsedDocument:
        """
        Parse un fichier DOCX en extrayant la hiérarchie des titres.
        Tente d'abord la détection par styles Word, puis par LLM.
        """
        doc = Document(file_path)
        paragraphs = doc.paragraphs
        total_paras = len([p for p in paragraphs if p.text.strip()])
        logger.info("Fichier chargé — %d paragraphes non vides", total_paras)

        # Stratégie 1 : détecter via les styles Heading natifs
        if self._has_heading_styles(paragraphs):
            logger.info("Stratégie : styles Heading natifs détectés")
            return self._parse_by_styles(paragraphs)

        # Stratégie 2 : détection par LLM à partir du formatage
        logger.info("Stratégie : détection des titres par LLM")
        return self._parse_by_llm(paragraphs)

    # =========================================================================
    # Stratégie 1 : Détection par styles Heading natifs
    # =========================================================================

    def _has_heading_styles(self, paragraphs) -> bool:
        """Vérifie si le document utilise au moins un style Heading."""
        return any(
            p.style.name.startswith('Heading')
            for p in paragraphs
            if p.style and p.style.name
        )

    def _parse_by_styles(self, paragraphs) -> ParsedDocument:
        """
        Extraction des titres via les styles Word natifs (Heading 1, Heading 2, etc.).
        C'est le cas le plus fiable quand le document est bien formaté.
        """
        result = ParsedDocument()
        sections = []
        current_content_lines = []

        for para in paragraphs:
            style_name = para.style.name if para.style else ''
            text = para.text.strip()
            if not text:
                continue

            if style_name.startswith('Heading'):
                if sections:
                    sections[-1].content = '\n'.join(current_content_lines)
                    current_content_lines = []

                try:
                    level = int(style_name.replace('Heading', '').strip())
                except ValueError:
                    level = 1

                sections.append(Section(title=text, level=level))
                if level == 1 and not result.title:
                    result.title = text
            else:
                current_content_lines.append(text)

        if sections and current_content_lines:
            sections[-1].content = '\n'.join(current_content_lines)

        result.sections = sections
        result.raw_text = '\n'.join(p.text for p in paragraphs if p.text.strip())
        return result

    # =========================================================================
    # Stratégie 2 : Détection par LLM
    # =========================================================================

    def _parse_by_llm(self, paragraphs) -> ParsedDocument:
        """
        Extraction des titres en envoyant les métadonnées de formatage
        de chaque paragraphe au LLM pour classification.
        """
        # Étape 0 : construire le mapping des numérotations automatiques Word
        doc_part = paragraphs[0].part if paragraphs else None
        num_formats = self._build_numbering_map(doc_part) if doc_part is not None else {}

        # Étape 1 : collecter les infos de formatage (avec numérotation)
        para_infos = self._collect_paragraph_info(paragraphs, num_formats)
        if not para_infos:
            return ParsedDocument()

        # Étape 2 : envoyer au LLM pour détecter les titres
        title_index, heading_levels = detect_headings_with_llm(para_infos)

        # Étape 3 : construire le document
        return self._build_document(para_infos, heading_levels, title_index)

    # =========================================================================
    # Extraction de la numérotation automatique Word
    # =========================================================================

    def _build_numbering_map(self, doc_part) -> dict:
        """
        Construit un mapping numId → {numFmt, lvlText, start} à partir
        des définitions de numérotation XML du document.

        Gère les chaînes numStyleLink → styleLink utilisées par certains
        documents Word (styles importés).
        """
        numbering_part = doc_part.numbering_part
        if numbering_part is None:
            return {}

        numbering_xml = numbering_part._element

        # Passe 1 : abstractNumId → format info (pour ceux qui ont des lvl directs)
        # + styleName → format info (pour ceux qui ont un styleLink)
        abstract_formats = {}
        style_link_formats = {}  # styleName → format info

        for abstract in numbering_xml.findall('.//w:abstractNum', NS):
            abs_id = abstract.get(f'{W}abstractNumId')

            # Vérifier si c'est un abstractNum avec styleLink (définition source)
            style_link_el = abstract.find('w:styleLink', NS)
            style_name = style_link_el.get(f'{W}val') if style_link_el is not None else None

            for lvl in abstract.findall('w:lvl', NS):
                ilvl = lvl.get(f'{W}ilvl')
                if ilvl != '0':
                    continue
                numFmt_el = lvl.find('w:numFmt', NS)
                lvlText_el = lvl.find('w:lvlText', NS)
                start_el = lvl.find('w:start', NS)
                fmt = {
                    'numFmt': numFmt_el.get(f'{W}val') if numFmt_el is not None else None,
                    'lvlText': lvlText_el.get(f'{W}val') if lvlText_el is not None else '',
                    'start': int(start_el.get(f'{W}val')) if start_el is not None else 1,
                }
                abstract_formats[abs_id] = fmt
                if style_name:
                    style_link_formats[style_name] = fmt

        # Passe 2 : résoudre les abstractNums avec numStyleLink
        # (qui référencent un style défini dans un autre abstractNum via styleLink)
        for abstract in numbering_xml.findall('.//w:abstractNum', NS):
            abs_id = abstract.get(f'{W}abstractNumId')
            if abs_id in abstract_formats:
                continue  # Déjà résolu

            num_style_link = abstract.find('w:numStyleLink', NS)
            if num_style_link is not None:
                linked_style = num_style_link.get(f'{W}val')
                if linked_style in style_link_formats:
                    abstract_formats[abs_id] = style_link_formats[linked_style]

        # Passe 3 : numId → format info (avec gestion des lvlOverride)
        num_formats = {}
        for num in numbering_xml.findall('.//w:num', NS):
            num_id = num.get(f'{W}numId')
            abstract_ref = num.find('w:abstractNumId', NS)
            if abstract_ref is None:
                continue

            abs_id = abstract_ref.get(f'{W}val')
            if abs_id not in abstract_formats:
                continue

            fmt = dict(abstract_formats[abs_id])  # Copie pour ne pas modifier l'original

            # Appliquer lvlOverride startOverride si présent
            for override in num.findall('w:lvlOverride', NS):
                ilvl = override.get(f'{W}ilvl')
                if ilvl != '0':
                    continue
                start_override = override.find('w:startOverride', NS)
                if start_override is not None:
                    fmt['start'] = int(start_override.get(f'{W}val'))

            num_formats[num_id] = fmt

        return num_formats

    def _get_numbering_prefix(self, para, num_formats: dict, counters: dict) -> str:
        """
        Retourne le préfixe de numérotation automatique d'un paragraphe.
        Ex: "A) ", "1. ", "a) ", "" (si pas de numérotation ou bullet).
        """
        numPr = para._p.find('.//w:numPr', NS)
        if numPr is None:
            return ''

        numId_el = numPr.find('w:numId', NS)
        if numId_el is None:
            return ''

        num_id = numId_el.get(f'{W}val')
        if num_id not in num_formats:
            return ''

        fmt_info = num_formats[num_id]
        numFmt = fmt_info['numFmt']

        # Ignorer les bullets (-, •, etc.)
        if numFmt == 'bullet':
            return ''

        # Incrémenter le compteur pour ce numId
        if num_id not in counters:
            counters[num_id] = fmt_info['start']
        else:
            counters[num_id] += 1

        counter_val = counters[num_id]
        lvlText = fmt_info['lvlText']

        # Convertir le compteur selon le format
        if numFmt == 'upperLetter':
            num_str = chr(64 + min(counter_val, 26))
        elif numFmt == 'lowerLetter':
            num_str = chr(96 + min(counter_val, 26))
        elif numFmt == 'upperRoman':
            num_str = self._to_roman(counter_val)
        elif numFmt == 'lowerRoman':
            num_str = self._to_roman(counter_val).lower()
        elif numFmt == 'decimal':
            num_str = str(counter_val)
        else:
            return ''

        # Appliquer le template lvlText (ex: "%1)" → "A)")
        prefix = lvlText.replace('%1', num_str)
        return prefix + ' '

    @staticmethod
    def _to_roman(n: int) -> str:
        """Convertit un entier en chiffres romains."""
        vals = [(1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'),
                (100, 'C'), (90, 'XC'), (50, 'L'), (40, 'XL'),
                (10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I')]
        result = ''
        for val, rom in vals:
            while n >= val:
                result += rom
                n -= val
        return result

    # =========================================================================
    # Collecte des métadonnées de formatage
    # =========================================================================

    def _collect_paragraph_info(self, paragraphs, num_formats: dict) -> list:
        """
        Collecte les informations de formatage de chaque paragraphe non vide.
        Préfixe le texte avec la numérotation automatique Word si présente.
        """
        para_infos = []
        counters = {}  # numId → compteur courant

        for para in paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Extraire la numérotation automatique Word (A), 1., a), etc.)
            prefix = self._get_numbering_prefix(para, num_formats, counters)
            if prefix:
                text = prefix + text

            runs_with_text = [r for r in para.runs if r.text.strip()]

            para_infos.append({
                'text': text,
                'font_size': self._get_font_size(runs_with_text),
                'is_bold': self._all_runs_bold(runs_with_text),
                'is_underline': self._majority_runs_underline(runs_with_text),
                'is_uppercase': text == text.upper() and any(c.isalpha() for c in text),
            })

        return para_infos

    def _get_font_size(self, runs: list) -> float:
        """Retourne la taille de police dominante des runs."""
        sizes = [r.font.size.pt for r in runs if r.font.size]
        if sizes:
            return max(set(sizes), key=sizes.count)
        return 0.0

    def _all_runs_bold(self, runs: list) -> bool:
        """Vérifie si TOUS les runs non vides sont en gras."""
        if not runs:
            return False
        return all(r.bold for r in runs)

    def _majority_runs_underline(self, runs: list) -> bool:
        """Vérifie si la majorité des runs (>50%) sont soulignés."""
        if not runs:
            return False
        underlined = sum(1 for r in runs if r.font.underline)
        return underlined > len(runs) / 2

    # =========================================================================
    # Construction du document final
    # =========================================================================

    def _build_document(
        self, para_infos: list, heading_levels: dict, title_index: int | None
    ) -> ParsedDocument:
        """
        Assemble le ParsedDocument final à partir des informations de paragraphes
        et des niveaux de titres retournés par le LLM.
        """
        result = ParsedDocument()
        sections = []
        current_content_lines = []

        if title_index is not None and title_index < len(para_infos):
            result.title = para_infos[title_index]['text']

        for i, info in enumerate(para_infos):
            text = info['text']
            level = heading_levels.get(i)

            if level is not None:
                if sections:
                    sections[-1].content = '\n'.join(current_content_lines)
                    current_content_lines = []

                sections.append(Section(title=text, level=level))

                if level == 1 and not result.title:
                    result.title = text
            else:
                current_content_lines.append(text)

        if sections and current_content_lines:
            sections[-1].content = '\n'.join(current_content_lines)

        if not result.title and para_infos:
            result.title = para_infos[0]['text']

        result.sections = sections
        result.raw_text = '\n'.join(p['text'] for p in para_infos)
        return result
